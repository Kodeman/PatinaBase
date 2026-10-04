"""Generate the golden fixture decks for the deck-import parser (SQ-356).

Run from the repository root (python-pptx 1.0.2 and Pillow 11):

    uv venv "$TMPDIR/deck-fixtures" --python python3
    uv pip install --python "$TMPDIR/deck-fixtures/bin/python" python-pptx==1.0.2 Pillow==11.3.0
    "$TMPDIR/deck-fixtures/bin/python" apps/designer-portal/src/lib/deck-import/__fixtures__/generate.py

Writes into this directory and overwrites the committed binaries:

    structure.pptx        reordered sldIdLst; nested, rotated and flipped groups; negative and positive
                          srcRect; picture-filled rounded rect; table with image fills; r:link pictures
                          (http and file:); SVG with PNG fallback; .wdp HD Photo layer; mc:AlternateContent;
                          EMF/WMF/TIFF; auto alt text; slide background picture; placeholder pictures that
                          inherit geometry from the layout and from the master.
    tickets.pptx          "ticket" captions (beside, below, table row, grouped text box, shared group),
                          a ticket listing 2 URLs, an overlay link shape, a 2x3 caption grid, an ambiguous
                          caption, a numbered legend on the same slide, a numbered "shopping list" slide,
                          notes with URLs, licence/attribution and denied links, a picture-less slide.
    keynote-2x.ppsx       Keynote-like 2x slide size, saved as a slideshow (.ppsx) content type.
    bomb-ratio.pptx       valid deck plus an 8 MiB entry of zeros (ratio ~1000:1).
    bomb-declared.pptx    valid deck whose central directory declares 3.75 GiB for one entry.
    zip64-298.pptx        fflate #298 layout: 0xFFFFFFFF size and offset sentinels with ZIP64 extras and
                          no ZIP64 end-of-central-directory locator.
    zip64-offset-only.pptx  only offset sentinels with ZIP64 extras, no locator (the variant size caps miss).
    zip64-valid.pptx      the same offset-only layout with a valid ZIP64 EOCD record and locator (accepted).
    legacy.ppt            a CFB (binary PowerPoint) header.
    keynote.key           a Keynote-style zip (Index/Document.iwa).
    not-a-deck.docx       a zip whose [Content_Types].xml is wordprocessingml.
    truth.json            element keys and expected values the jest suite asserts against.

Every zip is re-packed with a fixed timestamp so reruns are byte-stable.
"""

import io
import json
import os
import struct
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
W169, H169 = 12192000, 6858000
FIXED_TIME = (1980, 1, 1, 0, 0, 0)
truth = {}


# ------------------------------------------------------------------ images


def quad_image(fmt="PNG", size=(100, 80), colors=None, alpha=False):
    """Four coloured quadrants, so a crop is visible in which colours survive."""
    colors = colors or [(220, 40, 40), (40, 160, 60), (40, 70, 200), (230, 200, 40)]
    w, h = size
    mode = "RGBA" if alpha else "RGB"
    im = Image.new(mode, size, (0, 0, 0, 0) if alpha else (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w // 2 - 1, h // 2 - 1), fill=colors[0])
    d.rectangle((w // 2, 0, w - 1, h // 2 - 1), fill=colors[1])
    d.rectangle((0, h // 2, w // 2 - 1, h - 1), fill=colors[2])
    d.rectangle((w // 2, h // 2, w - 1, h - 1), fill=colors[3])
    buf = io.BytesIO()
    im.save(buf, fmt)
    return buf.getvalue()


def solid(color, fmt="PNG", size=(40, 30)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, fmt)
    return buf.getvalue()


PALETTE = [(180, 60, 60), (60, 140, 90), (70, 90, 170), (200, 150, 60), (120, 70, 150), (60, 150, 160), (150, 150, 150), (90, 60, 40)]


def img(i, fmt="PNG"):
    return solid(PALETTE[i % len(PALETTE)], fmt, size=(40 + i, 30 + i))


# ------------------------------------------------------------------ helpers


def key(slide, shape_id):
    return f"{slide.part.partname.lstrip('/')}#{shape_id}"


def next_id(slide):
    return slide.shapes._next_shape_id


def add_raw_part(prs, name, content_type, blob):
    return Part(PackURI(name), content_type, prs.part.package, blob)


def insert_shape_xml(slide, xml):
    el = parse_xml(xml)
    slide.shapes._spTree.append(el)
    return el


def pic_xml(shape_id, name, blip_attrs, x, y, w, h, blip_inner="", descr=None, src_rect=None, nv_pr=""):
    descr_attr = f' descr="{descr}"' if descr is not None else ""
    src = ""
    if src_rect:
        l, t, r, b = src_rect
        src = f'<a:srcRect l="{l}" t="{t}" r="{r}" b="{b}"/>'
    return (
        f'<p:pic {nsdecls("p", "a", "r")}><p:nvPicPr><p:cNvPr id="{shape_id}" name="{name}"{descr_attr}/>'
        f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr>{nv_pr}</p:nvPr></p:nvPicPr>'
        f'<p:blipFill><a:blip {blip_attrs}>{blip_inner}</a:blip>{src}<a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
    )


def set_group_xfrm(grp, off, ext, ch_off, ch_ext, rot=0, flip_h=False, flip_v=False):
    xfrm = grp._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    xfrm.find(qn("a:off")).set("x", str(off[0]))
    xfrm.find(qn("a:off")).set("y", str(off[1]))
    xfrm.find(qn("a:ext")).set("cx", str(ext[0]))
    xfrm.find(qn("a:ext")).set("cy", str(ext[1]))
    xfrm.find(qn("a:chOff")).set("x", str(ch_off[0]))
    xfrm.find(qn("a:chOff")).set("y", str(ch_off[1]))
    xfrm.find(qn("a:chExt")).set("cx", str(ch_ext[0]))
    xfrm.find(qn("a:chExt")).set("cy", str(ch_ext[1]))
    if rot:
        xfrm.set("rot", str(rot))
    if flip_h:
        xfrm.set("flipH", "1")
    if flip_v:
        xfrm.set("flipV", "1")


def textbox(slide, x, y, w, h, text):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tb.text_frame.text = text
    return tb


def add_run_link(paragraph, text, url):
    r = paragraph.add_run()
    r.text = text
    r.hyperlink.address = url
    return r


def blip_fill_shape(slide, img_bytes, x, y, w, h, crop, geom=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(geom, Emu(x), Emu(y), Emu(w), Emu(h))
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(img_bytes))
    l, t, r, b = crop
    blip = parse_xml(
        f'<a:blipFill {nsdecls("a", "r")} rotWithShape="1"><a:blip r:embed="{rid}"/>'
        f'<a:srcRect l="{l}" t="{t}" r="{r}" b="{b}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
    )
    sp._element.spPr.find(qn("a:prstGeom")).addnext(blip)
    return sp


def cell_image_fill(slide, cell, img_bytes):
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(img_bytes))
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<a:blipFill {nsdecls("a", "r")}><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'))


def reorder(prs, order):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for i in order:
        lst.append(ids[i])


def repack(src_bytes, dst, content_type_rewrite=None, extra_entries=()):
    """Re-pack with a fixed timestamp (byte-stable), optionally rewriting [Content_Types].xml."""
    with zipfile.ZipFile(io.BytesIO(src_bytes)) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            if name == "[Content_Types].xml" and content_type_rewrite:
                data = content_type_rewrite(data)
            info = zipfile.ZipInfo(name, FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)
        for name, data in extra_entries:
            info = zipfile.ZipInfo(name, FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)


def save(prs, name, **kw):
    buf = io.BytesIO()
    prs.save(buf)
    path = os.path.join(HERE, name)
    repack(buf.getvalue(), path, **kw)
    return path


# ------------------------------------------------------------------ structure.pptx


def build_structure():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W169), Emu(H169)
    blank = prs.slide_layouts[6]
    t = {"elements": {}, "texts": {}, "skipped": {}}

    # slide1: groups -------------------------------------------------------
    s1 = prs.slides.add_slide(blank)
    g1 = s1.shapes.add_group_shape()
    g2 = g1.shapes.add_group_shape()
    pa = g2.shapes.add_picture(io.BytesIO(img(0)), Emu(0), Emu(0), Emu(2000000), Emu(1000000))
    set_group_xfrm(g2, (0, 0), (4000000, 4000000), (0, 0), (4000000, 4000000), rot=5400000)
    set_group_xfrm(g1, (1000000, 1000000), (4000000, 2000000), (0, 0), (8000000, 4000000))
    g3 = s1.shapes.add_group_shape()
    pb = g3.shapes.add_picture(io.BytesIO(img(1)), Emu(6500000), Emu(1500000), Emu(1000000), Emu(1000000))
    pb.rotation = 30
    set_group_xfrm(g3, (6000000, 1000000), (4000000, 3000000), (6000000, 1000000), (4000000, 3000000), flip_h=True)
    # a plain picture that is itself flipped vertically
    pc = s1.shapes.add_picture(io.BytesIO(img(2)), Emu(500000), Emu(4500000), Emu(1500000), Emu(1200000))
    pc._element.spPr.find(qn("a:xfrm")).set("flipV", "1")
    t["elements"]["nested_rot"] = dict(key=key(s1, pa.shape_id), groups=[g1.shape_id, g2.shape_id])
    t["elements"]["flipped_group"] = dict(key=key(s1, pb.shape_id))
    t["elements"]["flip_v"] = dict(key=key(s1, pc.shape_id))

    # slide2: crops, picture fill, table fills -------------------------------
    s2 = prs.slides.add_slide(blank)
    neg = s2.shapes.add_picture(io.BytesIO(quad_image()), Emu(500000), Emu(800000), Emu(2400000), Emu(1600000))
    neg.crop_left, neg.crop_right = -0.1, -0.1
    pos = s2.shapes.add_picture(io.BytesIO(quad_image(colors=[(10, 10, 10), (90, 90, 90), (160, 160, 160), (240, 240, 240)])), Emu(3200000), Emu(800000), Emu(2400000), Emu(1600000))
    pos.crop_left, pos.crop_top, pos.crop_bottom = 0.25, 0.1, 0.2
    fill = blip_fill_shape(s2, img(3), 6000000, 800000, 2000000, 2000000, (10000, 0, 10000, 0))
    fill.click_action.hyperlink.address = "https://www.hay.dk/products/rounded-tray"
    tbl_shape = s2.shapes.add_table(2, 3, Emu(500000), Emu(3500000), Emu(9000000), Emu(2000000))
    tbl = tbl_shape.table
    for c, wdt in enumerate((2000000, 4000000, 3000000)):
        tbl.columns[c].width = Emu(wdt)
    for r in range(2):
        tbl.rows[r].height = Emu(1000000)
    cell_image_fill(s2, tbl.cell(0, 0), img(4))
    tbl.cell(0, 1).text = "Oak Stool\nWorkshop Co.\n$450"
    add_run_link(tbl.cell(0, 2).text_frame.paragraphs[0], "Shop the stool", "https://www.hay.dk/stool")
    cell_image_fill(s2, tbl.cell(1, 0), img(5))
    tbl.cell(1, 1).text = "Linen Lamp $1,200"
    tbl.cell(1, 2).text = "https://example-lamps.com/linen"
    t["elements"]["neg_crop"] = dict(key=key(s2, neg.shape_id), src_rect=[-0.1, 0, -0.1, 0])
    t["elements"]["pos_crop"] = dict(key=key(s2, pos.shape_id), src_rect=[0.25, 0.1, 0, 0.2])
    t["elements"]["picture_fill"] = dict(key=key(s2, fill.shape_id))
    frame = key(s2, tbl_shape.shape_id)
    t["elements"]["table_r0"] = dict(key=f"{frame}/r0c0")
    t["elements"]["table_r1"] = dict(key=f"{frame}/r1c0")

    # slide3: media kinds ------------------------------------------------------
    s3 = prs.slides.add_slide(blank)
    sid = lambda: next_id(s3)  # noqa: E731
    # mc:AlternateContent: Choice and Fallback each hold a picture; only Choice counts
    _, rid_choice = s3.part.get_or_add_image_part(io.BytesIO(img(6)))
    _, rid_fallback = s3.part.get_or_add_image_part(io.BytesIO(img(7)))
    choice_id = sid()
    alt = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main">'
        '<mc:Choice Requires="a14">'
        + pic_xml(choice_id, "Choice Picture", f'r:embed="{rid_choice}"', 300000, 300000, 1200000, 900000)
        + "</mc:Choice><mc:Fallback>"
        + pic_xml(choice_id, "Fallback Picture", f'r:embed="{rid_fallback}"', 300000, 300000, 1200000, 900000)
        + "</mc:Fallback></mc:AlternateContent>"
    )
    insert_shape_xml(s3, alt)
    t["elements"]["alternate_choice"] = dict(key=key(s3, choice_id))
    # SVG with PNG fallback
    _, rid_png = s3.part.get_or_add_image_part(io.BytesIO(img(1)))
    svg_part = add_raw_part(prs, "/ppt/media/vector1.svg", "image/svg+xml", b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>')
    rid_svg = s3.part.relate_to(svg_part, RT.IMAGE)
    svg_id = sid()
    insert_shape_xml(
        s3,
        pic_xml(
            svg_id, "Vector", f'r:embed="{rid_png}"', 1800000, 300000, 1200000, 900000,
            blip_inner=(
                '<a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
                f'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="{rid_svg}"/>'
                "</a:ext></a:extLst>"
            ),
        ),
    )
    t["elements"]["svg_fallback"] = dict(key=key(s3, svg_id))
    # HD Photo (.wdp) layer: ignored, the main blip wins
    _, rid_main = s3.part.get_or_add_image_part(io.BytesIO(img(2, "JPEG")))
    wdp_part = add_raw_part(prs, "/ppt/media/hdphoto1.wdp", "image/vnd.ms-photo", b"II\xbc\x01wdp")
    rid_wdp = s3.part.relate_to(wdp_part, "http://schemas.microsoft.com/office/2007/relationships/hdphoto")
    wdp_id = sid()
    insert_shape_xml(
        s3,
        pic_xml(
            wdp_id, "HD Photo", f'r:embed="{rid_main}"', 3300000, 300000, 1200000, 900000,
            blip_inner=(
                '<a:extLst><a:ext uri="{BEBA8EAE-BF5A-486C-A8C5-ECC9F3942E4B}">'
                f'<a14:imgProps xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"><a14:imgLayer r:embed="{rid_wdp}"/></a14:imgProps>'
                "</a:ext></a:extLst>"
            ),
            descr="A wooden chair with a woven seat&#10;&#10;Description automatically generated",
        ),
    )
    t["elements"]["wdp_auto_alt"] = dict(key=key(s3, wdp_id))
    # r:link pictures: http kept as remote_url, file: reported unsupported
    rid_http = s3.part.relate_to("https://cdn.example.com/linked-sofa.jpg", RT.IMAGE, is_external=True)
    http_id = sid()
    insert_shape_xml(s3, pic_xml(http_id, "Linked http", f'r:link="{rid_http}"', 4800000, 300000, 1200000, 900000))
    rid_file = s3.part.relate_to("file:///C:/Users/leah/Pictures/sofa.jpg", RT.IMAGE, is_external=True)
    file_id = sid()
    insert_shape_xml(s3, pic_xml(file_id, "Linked file", f'r:link="{rid_file}"', 6300000, 300000, 1200000, 900000))
    t["elements"]["remote_http"] = dict(key=key(s3, http_id), remote_url="https://cdn.example.com/linked-sofa.jpg")
    t["skipped"]["remote_file"] = key(s3, file_id)
    # EMF, WMF, TIFF
    for label, partname, ctype, blob, x in (
        ("emf", "/ppt/media/image90.emf", "image/x-emf", b"\x01\x00\x00\x00emf", 300000),
        ("wmf", "/ppt/media/image91.wmf", "image/x-wmf", b"\xd7\xcd\xc6\x9awmf", 1800000),
    ):
        part = add_raw_part(prs, partname, ctype, blob)
        rid = s3.part.relate_to(part, RT.IMAGE)
        shape_id = sid()
        insert_shape_xml(s3, pic_xml(shape_id, label, f'r:embed="{rid}"', x, 1800000, 1200000, 900000))
        t["skipped"][label] = key(s3, shape_id)
    tiff = s3.shapes.add_picture(io.BytesIO(img(3, "TIFF")), Emu(3300000), Emu(1800000), Emu(1200000), Emu(900000))
    t["skipped"]["tiff"] = key(s3, tiff.shape_id)
    # slide background picture
    _, rid_bg = s3.part.get_or_add_image_part(io.BytesIO(img(5, "JPEG")))
    cSld = s3._element.find(qn("p:cSld"))
    cSld.insert(0, parse_xml(
        f'<p:bg {nsdecls("p", "a", "r")}><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid_bg}"/>'
        f'<a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'
    ))
    t["elements"]["background"] = dict(key=f"{s3.part.partname.lstrip('/')}#bg")

    # slide4: placeholder picture inheriting from the layout ----------------
    s4 = prs.slides.add_slide(prs.slide_layouts[8])
    s4.shapes.title.text = "Placeholder Room"
    pic_ph = [p for p in s4.placeholders if p.placeholder_format.idx == 1][0]
    placed = pic_ph.insert_picture(io.BytesIO(quad_image(size=(120, 60))))
    t["elements"]["layout_placeholder"] = dict(
        key=key(s4, placed.shape_id),
        bbox=[placed.left, placed.top, placed.width, placed.height],
        src_rect=[placed.crop_left, placed.crop_top, placed.crop_right, placed.crop_bottom],
    )
    t["slide4_title"] = "Placeholder Room"

    # slide5: object placeholder whose layout has no xfrm -> master body ------
    s5 = prs.slides.add_slide(prs.slide_layouts[1])
    s5.shapes.title.text = "Master Inherited"
    for ph in list(s5.placeholders):
        if ph.placeholder_format.idx == 1:
            ph._element.getparent().remove(ph._element)
    _, rid_m = s5.part.get_or_add_image_part(io.BytesIO(img(7)))
    m_id = next_id(s5)
    el = parse_xml(
        f'<p:pic {nsdecls("p", "a", "r")}><p:nvPicPr><p:cNvPr id="{m_id}" name="Content Placeholder"/>'
        f'<p:cNvPicPr><a:picLocks noGrp="1" noChangeAspect="1"/></p:cNvPicPr><p:nvPr><p:ph idx="1"/></p:nvPr></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{rid_m}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr/></p:pic>'
    )
    s5.shapes._spTree.append(el)
    master_body = [p for p in prs.slide_master.placeholders if p.placeholder_format.idx == 1][0]
    t["elements"]["master_placeholder"] = dict(
        key=key(s5, m_id), bbox=[master_body.left, master_body.top, master_body.width, master_body.height]
    )

    slides = [s1, s2, s3, s4, s5]
    order = [2, 0, 4, 1, 3]
    reorder(prs, order)
    t["order"] = [slides[i].part.partname.lstrip("/") for i in order]
    t["slide_size"] = [W169, H169]
    save(prs, "structure.pptx")
    truth["structure"] = t


# ------------------------------------------------------------------ tickets.pptx


def build_tickets():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W169), Emu(H169)
    blank = prs.slide_layouts[6]
    t = {"elements": {}, "texts": {}, "slides": {}}

    def pic(slide, i, x, y, w, h, label):
        p = slide.shapes.add_picture(io.BytesIO(img(i)), Emu(x), Emu(y), Emu(w), Emu(h))
        t["elements"][label] = key(slide, p.shape_id)
        return p

    # T1: ticket beside and ticket below
    s = prs.slides.add_slide(blank)
    t["slides"]["beside_below"] = s.part.partname.lstrip("/")
    pic(s, 0, 500000, 1500000, 3000000, 2500000, "sofa")
    textbox(s, 3700000, 1700000, 2400000, 2000000, 'Harbor Sofa\nMaker: Lawson-Fenning\n$4,200\n84"W x 38"D x 30"H\nhttps://www.rh.com/harbor-sofa')
    pic(s, 1, 7000000, 1000000, 3000000, 2500000, "cane_chair")
    tb = textbox(s, 7000000, 3650000, 3000000, 1200000, "Cane Chair\nSKU: CC-1042\n$1,150.00")
    add_run_link(tb.text_frame.add_paragraph(), "Shop", "https://www.cb2.com/cane-chair")

    # T2: ticket as a table row; ticket listing two URLs
    s = prs.slides.add_slide(blank)
    t["slides"]["table_two_urls"] = s.part.partname.lstrip("/")
    pic(s, 2, 500000, 800000, 2700000, 2000000, "console")
    row = s.shapes.add_table(1, 3, Emu(500000), Emu(2950000), Emu(2700000), Emu(450000)).table
    for c, wdt in enumerate((1200000, 700000, 800000)):
        row.columns[c].width = Emu(wdt)
    row.rows[0].height = Emu(450000)
    row.cell(0, 0).text = "Walnut Console"
    row.cell(0, 1).text = "$3,400"
    add_run_link(row.cell(0, 2).text_frame.paragraphs[0], "studiodunn.com", "https://studiodunn.com/console")
    pic(s, 3, 6500000, 800000, 2500000, 2500000, "sconce")
    textbox(s, 9200000, 900000, 2600000, 1800000, "Brass Sconce (pair)\nhttps://www.lumens.com/brass-sconce\nalso: bit.ly/3xSconce")

    # T3: overlay link, grouped text ticket, shared group
    s = prs.slides.add_slide(blank)
    t["slides"]["overlay_groups"] = s.part.partname.lstrip("/")
    pic(s, 4, 400000, 600000, 3000000, 2400000, "rug")
    ov = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(500000), Emu(700000), Emu(2600000), Emu(2000000))
    ov.fill.background()
    ov.line.fill.background()
    ov.click_action.hyperlink.address = "https://www.westelm.com/overlay-rug"
    pic(s, 5, 4300000, 600000, 2600000, 2400000, "bench")
    tg = s.shapes.add_group_shape()
    a = tg.shapes.add_textbox(Emu(7100000), Emu(800000), Emu(2000000), Emu(500000))
    a.text_frame.text = "Rattan Bench"
    b = tg.shapes.add_textbox(Emu(7100000), Emu(1350000), Emu(2000000), Emu(500000))
    b.text_frame.text = "$890"
    t["texts"]["grouped_ticket_group"] = tg.shape_id
    gg = s.shapes.add_group_shape()
    pouf = gg.shapes.add_picture(io.BytesIO(img(6)), Emu(400000), Emu(3800000), Emu(1800000), Emu(1600000))
    t["elements"]["pouf"] = key(s, pouf.shape_id)
    gt = gg.shapes.add_textbox(Emu(9500000), Emu(5800000), Emu(2200000), Emu(600000))
    gt.text_frame.text = "Moss Pouf $380"

    # T4: 2x3 caption grid
    s = prs.slides.add_slide(blank)
    t["slides"]["grid"] = s.part.partname.lstrip("/")
    for r in range(2):
        for c in range(3):
            k = r * 3 + c
            x, y = 600000 + c * 3700000, 400000 + r * 3100000
            pic(s, k, x, y, 3200000, 2200000, f"grid_{k}")
            textbox(s, x, y + 2300000, 3200000, 500000, f"Grid Item {k + 1}\n${k + 1}00")

    # T5: one caption equidistant between two pictures
    s = prs.slides.add_slide(blank)
    t["slides"]["ambiguous"] = s.part.partname.lstrip("/")
    pic(s, 0, 500000, 1500000, 3000000, 2500000, "amb_left")
    pic(s, 1, 5500000, 1500000, 3000000, 2500000, "amb_right")
    textbox(s, 3750000, 2000000, 1500000, 1200000, "Side Table\nhttps://www.article.com/side-table")

    # T6: numbered legend on the same slide
    s = prs.slides.add_slide(blank)
    t["slides"]["legend_same"] = s.part.partname.lstrip("/")
    for n in range(3):
        x = 400000 + n * 2700000
        pic(s, n + 2, x, 800000, 2400000, 2400000, f"legend_{n + 1}")
        textbox(s, x + 100000, 900000, 400000, 400000, f"{n + 1}")
    textbox(
        s, 8700000, 3700000, 3200000, 2600000,
        "Sources\n1. Arc Floor Lamp https://www.flos.com/arc\n2. Linen Sofa $5,100 https://www.article.com/linen-sofa\n3. Jute Rug https://www.wayfair.com/jute",
    )

    # T7 + T8: numbered pictures, shopping list on the next slide
    s = prs.slides.add_slide(blank)
    t["slides"]["numbered"] = s.part.partname.lstrip("/")
    for n in range(2):
        x = 1000000 + n * 5000000
        pic(s, n + 5, x, 1000000, 4000000, 3000000, f"shop_{n + 1}")
        textbox(s, x + 3500000, 4100000, 400000, 400000, f"{n + 1}")
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Shopping List"
    t["slides"]["shopping_list"] = s.part.partname.lstrip("/")
    textbox(s, 800000, 1800000, 10000000, 2000000, "1 Travertine Side Table https://www.rejuvenation.com/side-table\n2 Wool Throw https://www.schoolhouse.com/throw")

    # T9: notes URLs, attribution, denied links
    s = prs.slides.add_slide(blank)
    t["slides"]["notes_denied"] = s.part.partname.lstrip("/")
    pic(s, 7, 4000000, 1000000, 4000000, 3000000, "h_chair")
    attr = s.shapes.add_textbox(Emu(4000000), Emu(4100000), Emu(4000000), Emu(300000))
    p = attr.text_frame.paragraphs[0]
    add_run_link(p, "This Photo", "https://somebody.blog/photo")
    r = p.add_run()
    r.text = " by Unknown Author is licensed under "
    add_run_link(p, "CC BY-SA", "https://creativecommons.org/licenses/by-sa/3.0/")
    misc = s.shapes.add_textbox(Emu(300000), Emu(5600000), Emu(3000000), Emu(900000))
    p = misc.text_frame.paragraphs[0]
    add_run_link(p, "Next", "https://www.example.com/never")
    p.runs[0]._r.find(qn("a:rPr")).find(qn("a:hlinkClick")).set("action", "ppaction://hlinkshowjump?jump=nextslide")
    p2 = misc.text_frame.add_paragraph()
    add_run_link(p2, "Email us", "mailto:studio@example.com")
    p3 = misc.text_frame.add_paragraph()
    add_run_link(p3, "Photo", "https://unsplash.com/photos/abc")
    hover = s.shapes.add_textbox(Emu(9000000), Emu(5600000), Emu(2000000), Emu(400000))
    hover.text_frame.text = "Hover me"
    rid_hover = s.part.relate_to("https://www.example.com/hover-only", RT.HYPERLINK, is_external=True)
    hover._element.nvSpPr.cNvPr.append(parse_xml(f'<a:hlinkHover {nsdecls("a", "r")} r:id="{rid_hover}"/>'))
    s.notes_slide.notes_text_frame.text = "Source https://www.burkedecor.com/h-chair and see www.chairish.com/item/99"

    # T10: picture-less slide with a link and notes
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Inspiration"
    t["slides"]["no_pictures"] = s.part.partname.lstrip("/")
    textbox(s, 800000, 1800000, 8000000, 800000, "Mood: https://www.pinterest.com/board/1")
    s.notes_slide.notes_text_frame.text = "More at https://www.1stdibs.com/x"

    save(prs, "tickets.pptx")
    truth["tickets"] = t


# ------------------------------------------------------------------ keynote-2x.ppsx


def build_keynote():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W169 * 2), Emu(H169 * 2)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = s.shapes.add_picture(io.BytesIO(img(3)), Emu(2000000), Emu(2000000), Emu(8000000), Emu(6000000))
    p.click_action.hyperlink.address = "https://www.dwr.com/eames"
    p.rotation = 15
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    p2 = s2.shapes.add_picture(io.BytesIO(img(4)), Emu(12000000), Emu(1000000), Emu(6000000), Emu(4000000))

    def to_slideshow(data):
        return data.replace(
            b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
            b"application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
        )

    buf = io.BytesIO()
    prs.save(buf)
    base = buf.getvalue()
    repack(base, os.path.join(HERE, "keynote-2x.ppsx"), content_type_rewrite=to_slideshow)
    truth["keynote"] = dict(
        slide_size=[W169 * 2, H169 * 2],
        eames=key(s, p.shape_id),
        second=key(s2, p2.shape_id),
    )
    return base


# ------------------------------------------------------------------ hostile packages


def zip64_repack(src_bytes, dst, limit, strip_locator):
    """CPython writes 0xFFFFFFFF sentinels plus ZIP64 extras for values over ZIP64_LIMIT."""
    saved = zipfile.ZIP64_LIMIT
    zipfile.ZIP64_LIMIT = limit
    try:
        with zipfile.ZipFile(io.BytesIO(src_bytes)) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                out = zipfile.ZipInfo(info.filename, FIXED_TIME)
                out.compress_type = zipfile.ZIP_DEFLATED
                zout.writestr(out, zin.read(info.filename))
    finally:
        zipfile.ZIP64_LIMIT = saved
    with open(dst, "rb") as f:
        data = f.read()
    eocd = data.rfind(b"PK\x05\x06")
    loc = eocd - 20
    assert data[loc:loc + 4] == b"PK\x06\x07", "expected a ZIP64 EOCD locator"
    if not strip_locator:
        return
    z64_off = struct.unpack("<Q", data[loc + 8:loc + 16])[0]
    (_, _, _, _, _, _, n_disk, n_total, cd_size, cd_off) = struct.unpack("<4sQHHIIQQQQ", data[z64_off:z64_off + 56])
    plain = struct.pack("<4sHHHHIIH", b"PK\x05\x06", 0, 0, n_disk, n_total, cd_size, min(cd_off, 0xFFFFFFFF), 0)
    with open(dst, "wb") as f:
        f.write(data[:z64_off] + plain)


def build_hostile(base):
    # ratio bomb: 8 MiB of zeros deflates to ~8 KB
    repack(base, os.path.join(HERE, "bomb-ratio.pptx"), extra_entries=[("ppt/media/zeros.bin", bytes(8 * 1024 * 1024))])
    # declared-size bomb: patch one central-directory uncompressed size to 3.75 GiB
    path = os.path.join(HERE, "bomb-declared.pptx")
    repack(base, path)
    with open(path, "rb") as f:
        data = bytearray(f.read())
    cd = data.find(b"PK\x01\x02")
    struct.pack_into("<I", data, cd + 24, 0xF0000000)
    with open(path, "wb") as f:
        f.write(bytes(data))
    # #298: size sentinels (entries over 1 KiB) and offset sentinels, locator stripped
    zip64_repack(base, os.path.join(HERE, "zip64-298.pptx"), 1024, True)
    # offset-only sentinels: the limit sits above every entry size but below later offsets
    with zipfile.ZipFile(io.BytesIO(base)) as z:
        biggest = max(max(i.file_size, i.compress_size) for i in z.infolist())
    zip64_repack(base, os.path.join(HERE, "zip64-offset-only.pptx"), biggest + 1, True)
    zip64_repack(base, os.path.join(HERE, "zip64-valid.pptx"), biggest + 1, False)
    # legacy binary PowerPoint (CFB header) and a Keynote package
    with open(os.path.join(HERE, "legacy.ppt"), "wb") as f:
        f.write(bytes.fromhex("D0CF11E0A1B11AE1") + bytes(504))
    with zipfile.ZipFile(os.path.join(HERE, "keynote.key"), "w") as z:
        z.writestr(zipfile.ZipInfo("Index/Document.iwa", FIXED_TIME), b"\x00iwa")
    with zipfile.ZipFile(os.path.join(HERE, "not-a-deck.docx"), "w") as z:
        z.writestr(
            zipfile.ZipInfo("[Content_Types].xml", FIXED_TIME),
            b'<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            b'<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )


def main():
    build_structure()
    build_tickets()
    base = build_keynote()
    build_hostile(base)
    with open(os.path.join(HERE, "truth.json"), "w") as f:
        json.dump(truth, f, indent=2, sort_keys=True)
        f.write("\n")
    for name in sorted(os.listdir(HERE)):
        if name.endswith((".py", ".md")):
            continue
        print(f"{name}: {os.path.getsize(os.path.join(HERE, name))} bytes")


if __name__ == "__main__":
    main()
