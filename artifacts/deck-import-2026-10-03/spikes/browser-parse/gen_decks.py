"""Generate synthetic .pptx decks for the browser-parse spike (SQ-351).

Throwaway. Writes decks plus a ground-truth JSON per deck into OUT_DIR (never the repo).

Each deck has 40-120 slides of high-res JPEG/PNG "product" images with:
  - plain pictures, some with srcRect crops and hyperlinks on the picture,
  - picture-filled shapes (p:sp/p:spPr/a:blipFill) with crops and links,
  - groups whose xfrm scales children (chExt != ext), so group transforms matter,
  - speaker notes with URLs, media reused across slides,
  - a shuffled sldIdLst, so display order != slideN.xml file order.

Also writes two ZIP64 re-packs of the 100 MB deck for fflate issue #298:
  - deck-100-z64.pptx        entries with ZIP64 extras + a ZIP64 EOCD record and locator (valid ZIP64)
  - deck-100-z64-noloc.pptx  the #298 layout: 0xFFFFFFFF sentinels + ZIP64 extras, no ZIP64 EOCD

Usage: python gen_decks.py OUT_DIR [sizes...]   (sizes in MB, default 10 50 100 150)
"""

import io
import json
import os
import random
import struct
import sys
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

SLIDE_W, SLIDE_H = 12192000, 6858000  # 16:9 in EMU

# size MB -> slide count, image long edge, noise sigma
TIERS = {
    10: dict(slides=40, size=(1600, 1200), sigma=8),
    50: dict(slides=60, size=(2400, 1800), sigma=10),
    100: dict(slides=90, size=(3000, 2250), sigma=10),
    150: dict(slides=120, size=(3600, 2400), sigma=10),
}


class ImageFactory:
    """Unique images cropped from one big noisy canvas, so generation stays fast."""

    def __init__(self, size, sigma, rng):
        self.w, self.h = size
        self.rng = rng
        cw, ch = self.w * 2, self.h * 2
        noise = Image.merge("RGB", [Image.effect_noise((cw, ch), sigma) for _ in range(3)])
        g = Image.linear_gradient("L").resize((cw, ch))
        base = Image.merge("RGB", (g, g.rotate(90).resize((cw, ch)), Image.new("L", (cw, ch), 160)))
        self.canvas = Image.blend(base, noise, 0.35)

    def make(self, png=False):
        r = self.rng
        x, y = r.randrange(0, self.w), r.randrange(0, self.h)
        im = self.canvas.crop((x, y, x + self.w, y + self.h))
        d = ImageDraw.Draw(im)
        color = tuple(r.randrange(0, 256) for _ in range(3))
        cx, cy = r.randrange(self.w // 4, 3 * self.w // 4), r.randrange(self.h // 4, 3 * self.h // 4)
        rad = r.randrange(self.h // 8, self.h // 3)
        d.ellipse((cx - rad, cy - rad, cx + rad, cy + rad), fill=color)
        d.rectangle((cx - rad // 2, cy + rad // 2, cx + rad // 2, cy + rad), fill=color[::-1])
        buf = io.BytesIO()
        if png:
            # product cut-out: smaller RGBA PNG with a transparent surround
            im = im.resize((self.w // 2, self.h // 2)).convert("RGBA")
            mask = Image.new("L", im.size, 0)
            ImageDraw.Draw(mask).ellipse((im.width // 8, im.height // 8, 7 * im.width // 8, 7 * im.height // 8), fill=255)
            im.putalpha(mask)
            im.save(buf, "PNG")
        else:
            im.save(buf, "JPEG", quality=88)
        return buf.getvalue()


def emu_bbox(x, y, w, h):
    return [int(x), int(y), int(w), int(h)]


def set_link(shape, url):
    shape.click_action.hyperlink.address = url


def add_blip_fill_shape(slide, img_bytes, x, y, w, h, crop, geom):
    """Picture-filled autoshape: p:sp/p:spPr/a:blipFill (python-pptx has no API for it)."""
    sp = slide.shapes.add_shape(geom, Emu(x), Emu(y), Emu(w), Emu(h))
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(img_bytes))
    l, t, r, b = crop
    blip = parse_xml(
        f'<a:blipFill {nsdecls("a", "r")} rotWithShape="1"><a:blip r:embed="{rid}"/>'
        f'<a:srcRect l="{l}" t="{t}" r="{r}" b="{b}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
    )
    sp._element.spPr.find(qn("a:prstGeom")).addnext(blip)
    return sp


def crop_attrs(pic):
    return [round(pic.crop_left * 100000), round(pic.crop_top * 100000), round(pic.crop_right * 100000), round(pic.crop_bottom * 100000)]


def build_deck(mb, out_dir, seed=0):
    tier = TIERS[mb]
    rng = random.Random(seed * 1000 + mb)
    fac = ImageFactory(tier["size"], tier["sigma"], rng)
    target = mb * 1_000_000 * 0.97
    pool = []  # unique image bytes generated so far
    state = {"bytes": 0, "n": 0}

    def image():
        """A new unique image until the byte target is met, then reuse (dedupes to one media part)."""
        if state["bytes"] < target or not pool:
            state["n"] += 1
            data = fac.make(png=(state["n"] % 6 == 0))
            pool.append(data)
            state["bytes"] += len(data)
            return data
        return rng.choice(pool)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    truth = []

    def new_slide(key):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Emu(300000), Emu(150000), Emu(6000000), Emu(500000))
        tb.name = f"key-{key}"
        tb.text_frame.text = f"Slide {key}"
        return s

    key = 0
    while True:
        min_reached = key >= tier["slides"]
        if (min_reached and state["bytes"] >= target) or key >= 120:
            break
        s = new_slide(key)
        pics = []  # expected occurrences, in spTree order
        pattern = key % 5
        if key > 0 and key % 10 == 0 and pool:
            # reuse an earlier image on this slide (one media part, many references)
            data = rng.choice(pool)
            p = s.shapes.add_picture(io.BytesIO(data), Emu(8000000), Emu(4500000), Emu(2400000), Emu(1800000))
            pics.append(dict(kind="pic", bbox=emu_bbox(8000000, 4500000, 2400000, 1800000), crop=[0, 0, 0, 0], link=None))
        if pattern == 0:
            x, y, w, h = 600000, 900000, 7200000, 5400000
            p = s.shapes.add_picture(io.BytesIO(image()), Emu(x), Emu(y), Emu(w), Emu(h))
            p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = 0.12, 0.05, 0.08, 0.1
            url = f"https://vendor.example.com/sofa/{key}"
            set_link(p, url)
            pics.append(dict(kind="pic", bbox=emu_bbox(x, y, w, h), crop=crop_attrs(p), link=url))
            cap = s.shapes.add_textbox(Emu(8100000), Emu(900000), Emu(3500000), Emu(1200000))
            cap.text_frame.text = f"Harbor Sofa {key}\nVendor Co.\n$4,{key:03d}.00"
        elif pattern == 1:
            for i in range(3):
                x, y, w, h = 400000 + i * 3900000, 1500000, 3600000, 2700000
                p = s.shapes.add_picture(io.BytesIO(image()), Emu(x), Emu(y), Emu(w), Emu(h))
                url = None
                if i < 2:
                    url = f"https://shop.example.com/item/{key}-{i}"
                    set_link(p, url)
                pics.append(dict(kind="pic", bbox=emu_bbox(x, y, w, h), crop=[0, 0, 0, 0], link=url))
        elif pattern == 2:
            grp = s.shapes.add_group_shape()
            child = []
            for i in range(2):
                x, y, w, h = 1000000 + i * 4200000, 1200000, 4000000, 3000000
                p = grp.shapes.add_picture(io.BytesIO(image()), Emu(x), Emu(y), Emu(w), Emu(h))
                url = f"https://maker.example.com/chair/{key}-{i}" if i == 0 else None
                if url:
                    set_link(p, url)
                child.append((x, y, w, h, url))
            # shrink the group to half size and move it: children are in group space now
            xfrm = grp._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
            cox, coy = int(ch_off.get("x")), int(ch_off.get("y"))
            cex, cey = int(ch_ext.get("cx")), int(ch_ext.get("cy"))
            gx, gy, gcx, gcy = 2000000, 3000000, cex // 2, cey // 2
            off.set("x", str(gx)), off.set("y", str(gy))
            ext.set("cx", str(gcx)), ext.set("cy", str(gcy))
            sx, sy = gcx / cex, gcy / cey
            for (x, y, w, h, url) in child:
                pics.append(dict(kind="pic", bbox=emu_bbox(gx + (x - cox) * sx, gy + (y - coy) * sy, w * sx, h * sy), crop=[0, 0, 0, 0], link=url, grouped=True))
        elif pattern == 3:
            x, y, w, h = 700000, 1000000, 4500000, 4500000
            crop = [10000, 5000, 10000, 15000]
            sp = add_blip_fill_shape(s, image(), x, y, w, h, crop, MSO_SHAPE.OVAL if key % 2 else MSO_SHAPE.ROUNDED_RECTANGLE)
            url = f"https://studio.example.com/lamp/{key}"
            set_link(sp, url)
            pics.append(dict(kind="shapeFill", bbox=emu_bbox(x, y, w, h), crop=crop, link=url))
            x2, y2, w2, h2 = 6000000, 1000000, 5000000, 3750000
            p = s.shapes.add_picture(io.BytesIO(image()), Emu(x2), Emu(y2), Emu(w2), Emu(h2))
            pics.append(dict(kind="pic", bbox=emu_bbox(x2, y2, w2, h2), crop=[0, 0, 0, 0], link=None))
        else:
            x, y, w, h = 1500000, 800000, 8000000, 5600000
            p = s.shapes.add_picture(io.BytesIO(image()), Emu(x), Emu(y), Emu(w), Emu(h))
            p.crop_left, p.crop_right = 0.2, 0.2
            pics.append(dict(kind="pic", bbox=emu_bbox(x, y, w, h), crop=crop_attrs(p), link=None))
        notes = None
        if key % 3 == 0:
            notes = f"Source: https://notes.example.com/ref/{key} lead time 8 weeks"
            s.notes_slide.notes_text_frame.text = notes
        truth.append(dict(key=f"key-{key}", pictures=pics, notesUrl=bool(notes)))
        key += 1

    # shuffle display order: sldIdLst order no longer matches slideN.xml numbering
    order = list(range(len(truth)))
    rng.shuffle(order)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for i in order:
        lst.append(ids[i])
    truth = [truth[i] for i in order]

    path = os.path.join(out_dir, f"deck-{mb}.pptx")
    prs.save(path)
    meta = dict(
        deck=os.path.basename(path),
        bytes=os.path.getsize(path),
        slides=len(truth),
        uniqueImagesGenerated=len(pool),
        pictures=sum(len(t["pictures"]) for t in truth),
        slidesInDisplayOrder=truth,
    )
    with open(path + ".truth.json", "w") as f:
        json.dump(meta, f)
    return path, meta


def repack_zip64(src, dst, strip_locator):
    """Re-pack with ZIP64 extras on every entry over 1 MB.

    Lowering zipfile.ZIP64_LIMIT makes CPython write 0xFFFFFFFF sizes + a ZIP64 extra field
    (local and central headers) plus a ZIP64 EOCD record and locator. With strip_locator the
    ZIP64 EOCD record/locator are removed and a plain EOCD is written: the #298 layout.
    """
    saved = zipfile.ZIP64_LIMIT
    zipfile.ZIP64_LIMIT = 1 << 20
    try:
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                zout.writestr(info, zin.read(info.filename), compress_type=info.compress_type)
    finally:
        zipfile.ZIP64_LIMIT = saved
    if not strip_locator:
        return
    with open(dst, "rb") as f:
        data = f.read()
    eocd = data.rfind(b"PK\x05\x06")
    loc = eocd - 20
    assert data[loc:loc + 4] == b"PK\x06\x07", "expected ZIP64 EOCD locator before EOCD"
    z64_off = struct.unpack("<Q", data[loc + 8:loc + 16])[0]
    assert data[z64_off:z64_off + 4] == b"PK\x06\x06"
    (_, _, _, _, _, _, n_disk, n_total, cd_size, cd_off) = struct.unpack("<4sQHHIIQQQQ", data[z64_off:z64_off + 56])
    assert cd_off + cd_size == z64_off
    plain = struct.pack("<4sHHHHIIH", b"PK\x05\x06", 0, 0, n_disk, n_total, cd_size, cd_off, 0)
    with open(dst, "wb") as f:
        f.write(data[:z64_off] + plain)


def main():
    out_dir = sys.argv[1]
    sizes = [int(s) for s in sys.argv[2:]] or [10, 50, 100, 150]
    os.makedirs(out_dir, exist_ok=True)
    for mb in sizes:
        path, meta = build_deck(mb, out_dir)
        print(f"{meta['deck']}: {meta['bytes'] / 1e6:.1f} MB, {meta['slides']} slides, {meta['pictures']} picture occurrences, {meta['uniqueImagesGenerated']} unique images")
    if 100 in sizes:
        src = os.path.join(out_dir, "deck-100.pptx")
        for name, strip in (("deck-100-z64.pptx", False), ("deck-100-z64-noloc.pptx", True)):
            dst = os.path.join(out_dir, name)
            repack_zip64(src, dst, strip)
            with open(src + ".truth.json") as f:
                meta = json.load(f)
            meta["deck"] = name
            meta["bytes"] = os.path.getsize(dst)
            with open(dst + ".truth.json", "w") as f:
                json.dump(meta, f)
            print(f"{name}: {meta['bytes'] / 1e6:.1f} MB (ZIP64 re-pack, locator {'stripped' if strip else 'kept'})")


if __name__ == "__main__":
    main()
